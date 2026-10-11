#!/usr/bin/env python3
"""
Builds docs/talk/talk.pptx: the 18 slides of the talk as native PowerPoint (text boxes, shapes, a native line chart, tables),
speaker notes taken from the <aside> of slides/<id>.html. Coordinates are the slides' 1920x1080 px canvas (1 px = 6350 EMU).

  python build_pptx.py            # needs python-pptx; writes talk.pptx next to this file
"""
from __future__ import annotations

import html
import json
import os
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PX = 6350
SANS, MONO = "IBM Plex Sans", "JetBrains Mono"
DARK, LIGHT, LIGHT2, WHITE = "13202E", "F6F4EF", "ECE9E1", "FFFFFF"
ORANGE, ORANGE_D, ORANGE_L, ORANGE_S = "E0782F", "A8521A", "F0B488", "C9601C"
BLUE, BLUE_D, BLUE_L = "3B6FB6", "2C5592", "9DB7DD"
GREY, BORDER, MUTED, FOOT, ON_DARK, CARD_D = "B9C2CC", "DAD6CC", "4A5563", "6B7684", "BFC8D2", "1E3044"


def px(v):
    return Emu(int(v * PX))


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = px(1920), px(1080)
BLANK = prs.slide_layouts[6]


def notes_of(sid):
    t = open(os.path.join(HERE, "slides", f"{sid}.html"), encoding="utf-8").read()
    m = re.search(r"<aside>(.*?)</aside>", t, re.S)
    return html.unescape(m.group(1)).strip() if m else ""


def new_slide(sid, bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    s.notes_slide.notes_text_frame.text = notes_of(sid)
    return s


def rect(s, x, y, w, h, fill, line=None, radius=None, shape=None):
    sh = s.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE), px(x), px(y), px(w), px(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, paras, size=32, color=DARK, bold=False, font=SANS, align=None, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """paras: list of paragraphs; a paragraph is a str or a list of runs (text, {bold, color, font, size})."""
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        if align:
            p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for t, st in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = st.get("font", font)
            f.size = Pt(st.get("size", size) * 0.5)
            f.bold = st.get("bold", bold)
            f.color.rgb = rgb(st.get("color", color))
    return tb


def header(s, eyebrow, title, dark=False, lines=1):
    text(s, 128, 128, 1664, 34, [eyebrow], size=24, color=ORANGE if dark else ORANGE_D, font=MONO)
    size = 64 if len(title) <= 42 else 50
    text(s, 128, 176, 1664, 80 * lines, [title], size=size, bold=True, color=LIGHT if dark else DARK, spacing=1.05)


def footer(s, t, dark=False):
    text(s, 128, 980, 1664, 36, [t], size=24, color=ON_DARK if dark else FOOT)


def bar_row(s, y, label, width, color, value, label_w=520, h=52, bold=False, x0=128, size=28):
    text(s, x0, y, label_w, h, [label], size=size, bold=bold, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x0 + label_w + 24, y, width, h, color, radius=6)
    text(s, x0 + label_w + 24 + width + 20, y, 200, h, [value], size=30, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def table(s, x, y, w, rows, widths, size=30, header_fill=None, row_fills=None, aligns=None, row_h=64):
    shape = s.shapes.add_table(len(rows), len(rows[0]), px(x), px(y), px(w), px(row_h * len(rows)))
    t = shape.table
    t.first_row = True
    for j, fr in enumerate(widths):
        t.columns[j].width = px(w * fr)
    for i, row in enumerate(rows):
        t.rows[i].height = px(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb((row_fills or {}).get(i, header_fill if i == 0 and header_fill else (LIGHT if i % 2 else WHITE)))
            c.margin_left = c.margin_right = px(14)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.name, r.font.size, r.font.bold = SANS, Pt(size * 0.5), i == 0
            r.font.color.rgb = rgb(DARK)
            if aligns and aligns[j] == "r":
                p.alignment = PP_ALIGN.RIGHT
    return shape


# 1 cover -------------------------------------------------------------------------------------------------------------
s = new_slide("cover", DARK)
text(s, 128, 128, 1664, 40, ["IT-КОНФЕРЕНЦИЯ · 2026"], size=28, color=ORANGE, font=MONO)
text(s, 128, 300, 1664, 230, ["Своя модель автодополнения строк", "для плагинов IntelliJ"], size=80, bold=True, color=LIGHT, spacing=1.0)
text(s, 128, 640, 1664, 60, ["C# и Go · 50–100 млн параметров · CPU ноутбука · без внешних рантаймов"], size=40, color=ON_DARK)
rect(s, 128, 820, 1000, 84, CARD_D, radius=12)
text(s, 156, 820, 960, 84, [[("return Task.FromResult(", {"color": LIGHT}), ("new HelloReply { … });", {"color": "7F8C99"})]],
     size=32, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1180, 820, 612, 84, ["[Имя докладчика] · [компания]"], size=28, color=ON_DARK, anchor=MSO_ANCHOR.MIDDLE)

# 2 task --------------------------------------------------------------------------------------------------------------
s = new_slide("task", LIGHT)
header(s, "ЗАДАЧА", "Дописать строку серым текстом — локально")
rect(s, 128, 320, 804, 600, DARK, radius=20)
text(s, 176, 368, 720, 300, ["public async Task StopAsync()", "{", [("    await _client.", {"color": LIGHT}), ("StopAsync();", {"color": ORANGE})], "}"],
     size=30, font=MONO, color=ON_DARK, spacing=1.5)
text(s, 176, 720, 720, 150, ["Курсор — после точки. Серым — остаток строки, Tab принимает. Показываем, только если модель уверена."],
     size=26, color=ON_DARK, spacing=1.3)
for i, (h3, p) in enumerate([("CPU ноутбука", "строка ≤ 150 мс, холодный файл ≤ 1 с"),
                             ("Без внешних рантаймов", "свой движок: Kotlin + нативные int8-ядра, формат .cml"),
                             ("Маленькая модель", "50 МБ в плагине, 50–100 млн параметров")]):
    y = 320 + i * 208
    rect(s, 988, y, 804, 184, WHITE, line=BORDER, radius=16)
    text(s, 1020, y + 30, 740, 50, [h3], size=36, bold=True)
    text(s, 1020, y + 90, 740, 80, [p], size=28, color=MUTED)
footer(s, "2 · плагины idea-dotnet-support (C#) и idea-golang-support (Go)")

# 3 model -------------------------------------------------------------------------------------------------------------
s = new_slide("model", LIGHT)
header(s, "МОДЕЛЬ", "Маленький Llama: размер решает задержка")
text(s, 128, 330, 800, 520, ["Decoder-only трансформер: RMSNorm, RoPE, GQA, SwiGLU, общие эмбеддинги.", "Контекст 2048 токенов, BPE 16 384 на язык.",
                             "Промпт — fill-in-the-middle: модель видит код и до, и после курсора.", "Экспорт в int8 по столбцам — без потерь качества."],
     size=32, spacing=1.45)
table(s, 988, 330, 804, [["модель", "параметры", "строка, мс"], ["S · d512 × 8", "31 M", "46"], ["M · d640 × 10", "50 M", "67"], ["L · d768 × 12", "102 M", "122"]],
      [0.4, 0.3, 0.3], header_fill=LIGHT2, row_fills={2: "FBE3D0"}, aligns="lrr")
text(s, 988, 620, 804, 120, ["Нативные q8-ядра, 8 потоков, промпт 512 + 20 токенов. Бюджет — 150 мс: 50M с запасом, 100M на границе."],
     size=26, color=MUTED, spacing=1.3)
footer(s, "3 · замеры движка idea-ml-completion (NnBench, EPYC)")

# 4 baseline ----------------------------------------------------------------------------------------------------------
s = new_slide("baseline", LIGHT)
header(s, "ОТПРАВНАЯ ТОЧКА", "Одна и та же модель: 51.6 % или 32.7 %?")
rect(s, 128, 330, 808, 360, WHITE, line=BORDER, radius=20)
text(s, 176, 370, 720, 130, ["51.6 %"], size=120, bold=True, color="7F8C99")
text(s, 176, 520, 720, 50, ["тест движка"], size=36, bold=True)
text(s, 176, 580, 720, 90, ["3 000 позиций из тестовой части того же корпуса"], size=28, color=MUTED)
rect(s, 984, 330, 808, 360, DARK, radius=20)
text(s, 1032, 370, 720, 130, ["32.7 %"], size=120, bold=True, color=ORANGE)
text(s, 1032, 520, 720, 50, ["чистый набор"], size=36, bold=True, color=LIGHT)
text(s, 1032, 580, 720, 90, ["12 000 позиций из репозиториев, которых не видела ни одна модель"], size=28, color=ON_DARK)
text(s, 128, 740, 1664, 110, ["Строка угадана целиком — у готовой C#-модели плагина 50M. Первое правило: сравнивать только на одинаковом, "
                              "честно изолированном наборе."], size=34, spacing=1.3)
footer(s, "4 · cs-nn-50m-e3-lr2e3 · «строка целиком» = остаток строки совпал с точностью до хвостовых пробелов")

# 5 data --------------------------------------------------------------------------------------------------------------
s = new_slide("data", LIGHT)
header(s, "ДАННЫЕ", "Десятки тысяч репозиториев — без сборки кода")
for i, (cap, big, sub, dark) in enumerate([("C#: репозитории", "29 117", "6.47 млрд токенов", False),
                                           ("Go: репозитории", "43 356", "10.42 млрд токенов", False),
                                           ("позиции курсора", "1.05 млн", "на язык, ~1 млрд токенов", True)]):
    x = 128 + i * 566
    rect(s, x, 320, 534, 230, DARK if dark else WHITE, line=None if dark else BORDER, radius=16)
    text(s, x + 36, 350, 470, 34, [cap], size=24, color=ON_DARK if dark else MUTED)
    text(s, x + 36, 392, 470, 90, [big], size=72, bold=True, color=ORANGE if dark else DARK)
    text(s, x + 36, 486, 470, 40, [sub], size=28, color=ON_DARK if dark else MUTED)
for i, (h3, p) in enumerate([("Разбор, а не сборка", "Roslyn и go/types на снимке файла: ни restore, ни MSBuild, ни запуска чужого кода."),
                             ("Позиции по стратам", "после точки, в аргументах, после return, начало строки… Точная реконструкция до байта."),
                             ("Гигиена", "лицензии из allowlist, дедуп по хэшу, сгенерированный код и секреты — вон.")]):
    x = 128 + i * 568
    text(s, x, 610, 528, 46, [h3], size=32, bold=True)
    text(s, x, 664, 528, 200, [p], size=28, color=MUTED, spacing=1.3)
footer(s, "5 · публичные датасеты HF: csharp-ml-complation, go-ml-complation · два CPU-сервера по 128 ядер, ~1 сутки")

# 6 eval --------------------------------------------------------------------------------------------------------------
s = new_slide("eval", LIGHT2)
header(s, "ЧЕСТНАЯ ОЦЕНКА", "Проверка на коде, которого никто не видел")
text(s, 128, 330, 520, 230, ["0"], size=200, bold=True)
text(s, 128, 580, 520, 140, ["общих репозиториев у обучения и проверки — проверено скриптом"], size=32, color=MUTED, spacing=1.3)
for i, (n, t) in enumerate([("300", "свежих репозиториев на язык (5–19 ★), отложены до обучения"),
                            ("4–5 %", "файлов-почти-копий убраны из оценки (MinHash, Jaccard ≥ 0.8)"),
                            ("12 000", "позиций, движок плагина как декодер, его порог показа"),
                            ("95 %", "доверительный интервал: парный бутстрап на одних позициях")]):
    y = 340 + i * 130
    text(s, 696, y, 220, 60, [n], size=44, bold=True, color=ORANGE_D)
    text(s, 940, y + 4, 852, 110, [t], size=30, spacing=1.3)
footer(s, "6 · метрики: строка целиком · доля показанных подсказок · точность показанных")

# 7 context -----------------------------------------------------------------------------------------------------------
s = new_slide("context", LIGHT)
header(s, "ЭКСПЕРИМЕНТ 1 · КОНТЕКСТ", "Факты о типах не помогли. Задача — помогла")
for i, (lab, w, col, val, b) in enumerate([("Готовая модель плагина", 768, GREY, "32.9 %", False),
                                           ("Дообучение на задаче, без контекста", 923, ORANGE, "39.5 %", True),
                                           ("С контекстом: только профиль", 912, BLUE, "39.1 %", False),
                                           ("С контекстом: профиль + факты", 901, BLUE, "38.6 %", False)]):
    bar_row(s, 320 + i * 76, lab, w, col, val, bold=b)
text(s, 128, 660, 800, 150, [[("+6.6 п.п.", {"bold": True}), (" — само дообучение на позициях курсора (95 %: +5.8…+7.4).", {})]], size=30, spacing=1.3)
text(s, 968, 660, 824, 200, [[("−0.95 п.п.", {"bold": True}), (" — блок LOCAL / ARG / RECV / MEMBER: повторяет видимый код, а нужного члена "
                                                              "в обрезанном списке часто нет.", {})]], size=30, spacing=1.3)
footer(s, "7 · C# 50M, 311 тыс. позиций, 2× RTX PRO 6000, ~1 ч · строка целиком, 12 000 позиций")

# 8 format ------------------------------------------------------------------------------------------------------------
s = new_slide("format", LIGHT)
header(s, "ДООБУЧЕНИЕ НА ЗАДАЧЕ", "Учим ровно на том промпте, который соберёт плагин")
x = 128
for w, cap, lab, hot in [(362, "fim_prefix fim_suffix", "код после строки", False), (362, "fim_middle file_sep", "путь файла", False),
                         (470, "≤ 1024 токена", "код до курсора ⟨⟩", False), (470, "loss только здесь", "остаток строки + стоп", True)]:
    rect(s, x, 330, w, 170, "FBE3D0" if hot else "DCE5F1", line=ORANGE if hot else BLUE)
    text(s, x + 26, 360, w - 52, 34, [cap], size=24, color=ORANGE_D if hot else BLUE_D, font=MONO)
    text(s, x + 26, 410, w - 52, 80, [lab], size=30, bold=True)
    x += w
for i, (h3, p) in enumerate([("Байт в байт", "Сборка промпта совпала с движком плагина на 1 000 из 1 000 позиций, лечение токенов — 3 000 из 3 000."),
                             ("Ответ — одна строка", "Остаток физической строки и стоп-токен; переносов внутри ответа нет."),
                             ("Один проход", "1.05 млн документов, ~1 млрд токенов, 20–45 минут на 2× H200.")]):
    x = 128 + i * 568
    text(s, x, 570, 528, 46, [h3], size=32, bold=True)
    text(s, x, 626, 528, 200, [p], size=28, color=MUTED, spacing=1.3)
footer(s, "8 · SPM fill-in-the-middle, как у движка idea-ml-completion")

# 9 pretrain (native chart) -------------------------------------------------------------------------------------------
s = new_slide("pretrain", LIGHT)
header(s, "ОБУЧЕНИЕ ЯЗЫКУ С НУЛЯ", "100M стабильно впереди 50M на тех же данных")
cd = CategoryChartData()
grid = list(range(20, 101, 10))
cs_x = [100 * k / 12343 for k in (2000, 4000, 6000, 8000, 10000, 12000, 12343)]
go_x = [100 * k / 19876 for k in (2000, 4000, 6000, 8000, 10000, 12000, 14000, 16000, 18000, 19876)]


def interp(xs, ys, x):
    for (a, ya), (b, yb) in zip(zip(xs, ys), zip(xs[1:], ys[1:])):
        if a <= x <= b:
            return round(ya + (yb - ya) * (x - a) / (b - a), 3)
    return None


series = [("C# 50M · 3.39", cs_x, [4.65, 4.19, 3.75, 3.59, 3.46, 3.40, 3.39], ORANGE_L),
          ("C# 100M · 3.19", cs_x, [4.58, 4.07, 3.59, 3.41, 3.27, 3.20, 3.19], ORANGE_S),
          ("Go 50M · 2.87", go_x, [3.88, 3.48, 3.34, 3.21, 3.13, 3.05, 2.99, 2.93, 2.89, 2.87], BLUE_L),
          ("Go 100M · 2.67", go_x, [3.61, 3.24, 3.11, 3.00, 2.92, 2.84, 2.78, 2.72, 2.68, 2.67], BLUE_D)]
cd.categories = [f"{g} %" for g in grid]
for name, xs, ys, _ in series:
    cd.add_series(name, [interp(xs, ys, g) for g in grid])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, px(128), px(300), px(1664), px(640), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.RIGHT
ch.legend.include_in_layout = False
ch.legend.font.size, ch.legend.font.name = Pt(13), SANS
for i, (_, _, _, col) in enumerate(series):
    ln = ch.series[i].format.line
    ln.color.rgb = rgb(col)
    ln.width = Pt(3.5 if i % 2 else 2.5)
    ch.series[i].smooth = False
    ch.series[i].marker.style = XL_MARKER_STYLE.NONE
va, ca = ch.value_axis, ch.category_axis
va.minimum_scale, va.maximum_scale, va.major_unit = 2.5, 4.8, 0.5
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = rgb(BORDER)
for ax, title in ((va, "perplexity на проверке"), (ca, "доля обучения")):
    ax.tick_labels.font.size, ax.tick_labels.font.name = Pt(12), SANS
    ax.tick_labels.font.color.rgb = rgb(FOOT)
    ax.has_title = True
    ax.axis_title.text_frame.text = title
    ax.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
footer(s, "9 · lr 2e-3 → 2e-4, 0.5 млн токенов/шаг, FIM 0.7 · C# 50M 73 мин, 100M 2 ч 18 мин, Go 100M ~3 ч на 2× H200")

# 10 loss trap --------------------------------------------------------------------------------------------------------
s = new_slide("losstrap", DARK)
header(s, "ЛОВУШКА", "Loss дообучения рос — а модель становилась лучше", dark=True)
for i, (big, col, h3, p) in enumerate([("0.78 → 1.07", ORANGE, "loss на ответах, первая треть", "а строк целиком в итоге: 33.5 % → 40.5 %"),
                                       ("3.6", BLUE_L, "loss на стоп-токене", "модель на целых файлах не видела «конец ответа» после строки; "
                                                                                "выучивает его за ~100 шагов")]):
    x = 128 + i * 856
    rect(s, x, 320, 808, 380, CARD_D, radius=20)
    text(s, x + 44, 360, 720, 110, [big], size=96, bold=True, color=col)
    text(s, x + 44, 490, 720, 50, [h3], size=34, bold=True, color=LIGHT)
    text(s, x + 44, 550, 720, 130, [p], size=28, color=ON_DARK, spacing=1.3)
text(s, 128, 750, 1664, 120, ["Вывод: решения — только по качеству на чистых позициях. Замер каждые 4 000 шагов, ~1 минута на карте рядом с обучением."],
     size=34, color=LIGHT, spacing=1.3)
footer(s, "10 · гипотеза «модель запомнила файлы» отвергнута: loss на обучающих и невиденных позициях одинаков (1.19 / 1.19)", dark=True)

# 11 lr ---------------------------------------------------------------------------------------------------------------
s = new_slide("lr", LIGHT)
header(s, "ЭКСПЕРИМЕНТ 2 · ШАГ ДООБУЧЕНИЯ", "Маленький шаг обучения — главный рычаг")
for i, (lab, w, col, val, b) in enumerate([("lr 2e-4, весь объём", 871, GREY, "38.7 %", False), ("lr 1e-4, короткий (18 %)", 929, BLUE_L, "41.3 %", False),
                                           ("lr 2e-5, короткий (18 %)", 934, BLUE_L, "41.5 %", False), ("lr 5e-5, весь объём", 972, ORANGE, "43.2 %", True)]):
    bar_row(s, 320 + i * 80, lab, w, col, val, h=56, bold=b)
text(s, 128, 680, 800, 150, [[("На полном наборе: ", {}), ("40.5 % → 43.8 %", {"bold": True}), (" строк целиком, +3.3 п.п. от одной настройки.", {})]],
     size=30, spacing=1.3)
text(s, 968, 680, 824, 150, ["Короткие прогоны по 15 минут на запасной H100 — подсказка, полный прогон — ответ."], size=30, spacing=1.3)
footer(s, "11 · C# 50M, строка целиком на 2 000 позициях; 18 % = 6 000 шагов со своим косинусом")

# 12 results ----------------------------------------------------------------------------------------------------------
s = new_slide("results", LIGHT2)
header(s, "РЕЗУЛЬТАТ · 50M, ТА ЖЕ СКОРОСТЬ", "+11 п.п. для C#, +17 п.п. для Go")
for i, (lang, ow, nw, col, ov, nv, rest) in enumerate([("C#", 327, 438, ORANGE, "старая 32.7 %", "новая 43.8 %",
                                                        ["показ 15.4 → 28.9 % (×1.9)", "точность показанных 81.3 → 89.6 %"]),
                                                       ("Go", 348, 518, BLUE, "старая 34.8 %", "новая 51.8 %",
                                                        ["показ 21.5 → 37.2 % (×1.7)", "точность показанных 80.9 → 91.3 %"])]):
    x = 128 + i * 852
    rect(s, x, 320, 812, 420, LIGHT, line=BORDER, radius=20)
    text(s, x + 44, 352, 700, 54, [lang], size=40, bold=True)
    rect(s, x + 44, 430, ow, 40, GREY, radius=6)
    text(s, x + 60 + ow, 430, 260, 40, [ov], size=28, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x + 44, 490, nw, 40, col, radius=6)
    text(s, x + 60 + nw, 490, 260, 40, [nv], size=28, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 44, 570, 720, 110, rest, size=28, color=MUTED, spacing=1.4)
text(s, 128, 780, 1664, 110, ["Тот же размер, тот же движок, та же задержка: подсказок почти вдвое больше, а ошибок среди показанных — вдвое меньше."],
     size=32, spacing=1.3)
footer(s, "12 · строка целиком, 12 000 чистых позиций; C# +11.0 п.п. (95 %: +10.3…+11.8), Go +17.0 п.п. (+16.1…+17.8)")

# 13 kinds ------------------------------------------------------------------------------------------------------------
s = new_slide("kinds", LIGHT)
header(s, "ГДЕ ВЫИГРЫШ · C# 50M", "После ключевых слов и операторов — с нуля до половины", lines=2)
rows = [("после return / new", 43, 586, "4 → 49 %"), ("после оператора", 8, 572, "1 → 48 %"), ("лямбды", 18, 538, "2 → 45 %"),
        ("аргументы", 342, 577, "29 → 48 %"), ("LINQ", 347, 487, "29 → 41 %"), ("начало строки", 308, 386, "26 → 32 %"),
        ("после точки", 552, 588, "46 → 49 %"), ("недописанное имя", 631, 673, "53 → 56 %")]
for i, (lab, ow, nw, val) in enumerate(rows):
    y = 360 + i * 66
    text(s, 128, y, 330, 44, [lab], size=26, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 478, y, max(ow, 4), 18, GREY)
    rect(s, 478, y + 22, nw, 18, ORANGE)
    text(s, 478 + max(ow, nw) + 20, y, 200, 44, [val], size=26, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 1372, 370, 36, 18, GREY)
text(s, 1420, 362, 360, 34, ["старая модель"], size=26)
rect(s, 1372, 420, 36, 18, ORANGE)
text(s, 1420, 412, 360, 34, ["новая 50M"], size=26)
text(s, 1372, 480, 420, 360, ["Старая модель почти не умела начинать выражение. После точки и в именах прирост скромный: там нужны знания о типах из "
                              "других файлов."], size=28, color=MUTED, spacing=1.35)
footer(s, "13 · строка целиком по типам позиции курсора, 12 000 позиций, промпты прежнего образца")

# 14 size -------------------------------------------------------------------------------------------------------------
s = new_slide("size", LIGHT)
header(s, "РАЗМЕР · 100M", "Вдвое больше параметров — +2…4 п.п. за 2× задержки", lines=2)
y = 360
for lang, items in (("C#", [("старая", 472, GREY, "39.3 %", False), ("50M", 532, ORANGE_L, "44.3 %", False), ("100M", 556, ORANGE_S, "46.3 %", True)]),
                    ("Go", [("старая", 548, GREY, "45.7 %", False), ("50M", 630, BLUE_L, "52.5 %", False), ("100M", 672, BLUE_D, "56.0 %", True)])):
    text(s, 128, y, 300, 44, [lang], size=34, bold=True)
    y += 54
    for lab, w, col, val, b in items:
        bar_row(s, y, lab, w, col, val, label_w=150, h=40, bold=b, size=26)
        y += 52
    y += 14
table(s, 1272, 360, 520, [["цена", "50M", "100M"], ["строка", "67 мс", "122 мс"], ["файл .cml", "50 МБ", "103 МБ"], ["память", "~100 МБ", "~200 МБ"]],
      [0.42, 0.29, 0.29], size=28, header_fill=LIGHT2, aligns="lrr", row_h=60)
text(s, 1272, 640, 520, 200, ["Показ: C# 29.6 → 32.2 %, Go 39.0 → 44.0 % при той же точности. 100M — опция для мощных машин."], size=28,
     color=MUTED, spacing=1.35)
footer(s, "14 · строка целиком, 12 000 позиций, промпты нового образца; C# 100M − 50M: +2.0 п.п.; Go: +3.5 п.п.")

# 15 spaces -----------------------------------------------------------------------------------------------------------
s = new_slide("spaces", LIGHT2)
header(s, "БАГ ИЗ ПЛАГИНА", "newHelloReply: куда пропал пробел")
rect(s, 128, 310, 812, 330, DARK, radius=20)
text(s, 168, 350, 740, 120, [[("до обучения:  ", {}), ("new│␠Hel│lo│Reply", {"color": LIGHT})], [("дообученная: ", {}), ("new│Hello│Reply", {"color": ORANGE})]],
     size=28, font=MONO, color=ON_DARK, spacing=1.6)
text(s, 168, 500, 740, 110, ["Пробел — часть токена следующего слова. Модель выбрала токен без пробела."], size=26, color=ON_DARK, spacing=1.35)
text(s, 980, 310, 812, 44, ["Причина"], size=32, bold=True)
text(s, 980, 360, 812, 150, [[("После ", {}), ("return ⟨⟩", {"bold": True}), (" в редакторе ничего нет: промпт кончался токеном-пробелом, а ответ — "
                                                                             "словом без пробела. Так было в 14 % документов C# и 22 % Go.", {})]],
     size=28, color=MUTED, spacing=1.3)
text(s, 980, 520, 812, 44, ["Исправление"], size=32, bold=True)
text(s, 980, 570, 812, 80, ["Пробел перед курсором уходит в ответ — и в данных, и в движке плагина."], size=28, color=MUTED, spacing=1.3)
for i, (big, col, cap) in enumerate([("32.7 → 39.3 %", ORANGE_D, "старая C#-модель — только от нового правила промпта"),
                                     ("34.8 → 45.7 %", ORANGE_D, "старая Go-модель, без всякого переобучения"),
                                     ("16 → 5", BLUE_D, "потерянных пробелов на 2 000 ответов Go")]):
    x = 128 + i * 566
    text(s, x, 700, 534, 70, [big], size=56, bold=True, color=col)
    text(s, x, 780, 534, 80, [cap], size=26, color=MUTED, spacing=1.3)
footer(s, "15 · C#: редкие склейки посреди строки остались (2–3 из 2 000) — следующий шаг: подмешивать данные обучения языку")

# 16 summary ----------------------------------------------------------------------------------------------------------
s = new_slide("summary", LIGHT)
header(s, "СВОДКА", "Все модели на одном наборе")
table(s, 128, 320, 1664, [["модель", "строка целиком", "показ", "точность", ".cml"],
                          ["C# · старая в плагине", "39.3 %", "16.9 %", "87.7 %", "50 МБ"],
                          ["C# · 50M", "44.3 %", "29.6 %", "89.7 %", "50 МБ"],
                          ["C# · 100M", "46.3 %", "32.2 %", "89.9 %", "103 МБ"],
                          ["Go · старая в плагине", "45.7 %", "26.2 %", "86.2 %", "50 МБ"],
                          ["Go · 50M", "52.5 %", "39.0 %", "91.7 %", "50 МБ"],
                          ["Go · 100M", "56.0 %", "44.0 %", "92.3 %", "103 МБ"]],
      [0.34, 0.18, 0.16, 0.18, 0.14], size=32, header_fill=LIGHT2, row_fills={2: "FBE3D0", 5: "DCE5F1"}, aligns="lrrrr", row_h=70)
text(s, 128, 840, 1664, 60, ["Выделены модели, которые идут в плагины по умолчанию: тот же размер и скорость, что у старых."], size=30)
footer(s, "16 · 12 000 чистых позиций, промпты нового образца; точность = доля верных среди показанных плагином")

# 17 lessons ----------------------------------------------------------------------------------------------------------
s = new_slide("lessons", DARK)
header(s, "УРОКИ", "Шесть вещей, которые мы бы сделали сразу", dark=True)
cards = [("Учить на задаче", "Тот же промпт, что у плагина, loss на ответе: +7…+16 п.п.", ORANGE),
         ("Маленький шаг", "5e-5 вместо 2e-4 при дообучении: +3.3 п.п.", ORANGE),
         ("Не верить loss", "Мерить качество каждые 4 000 шагов на чистых позициях.", ORANGE),
         ("Чистый набор", "Свежие репозитории, ноль пересечений, парный бутстрап.", BLUE_L),
         ("Проверять идеи дёшево", "Контекст из IDE не помог — узнали за час, а не за месяц.", BLUE_L),
         ("Смотреть на токены", "Один пробел в промпте — до +11 п.п. у старой модели.", BLUE_L)]
for i, (h3, p, col) in enumerate(cards):
    x, y = 128 + (i % 3) * 564, 320 + (i // 3) * 300
    rect(s, x, y, 536, 272, CARD_D, radius=16)
    text(s, x + 36, y + 36, 464, 50, [h3], size=32, bold=True, color=col)
    text(s, x + 36, y + 100, 464, 150, [p], size=26, color=ON_DARK, spacing=1.35)
footer(s, "17 · и инфраструктура: снимки каждые 20 минут, монитор, перенос обучения между серверами без потерь", dark=True)

# 18 next -------------------------------------------------------------------------------------------------------------
s = new_slide("next", LIGHT)
header(s, "ЧТО ДАЛЬШЕ", "Следующие шаги")
items = [("Новое правило пробелов в движке", " обоих плагинов — бесплатный прирост и старым моделям."),
         ("Подмешивать данные обучения языку", " в дообучение — убрать редкие склейки в C#."),
         ("Многострочные подсказки", ": тело метода после «{» с отдельным стоп-токеном."),
         ("Знания из соседних файлов", " для позиций после точки — там прирост пока скромный."),
         ("100M как режим «качество»", " — после замеров задержки на ноутбуках.")]
for i, (b, rest) in enumerate(items):
    y = 320 + i * 84
    text(s, 128, y, 64, 50, [f"{i + 1:02d}"], size=32, color=ORANGE_D, font=MONO)
    text(s, 220, y, 1572, 60, [[(b, {"bold": True}), (rest, {})]], size=30)
text(s, 128, 780, 1664, 80, ["Данные: huggingface.co/datasets/dvislobokov/csharp-ml-complation · go-ml-complation · код: "
                             "github.com/dvislobokov/csharp-dataset-prepare"], size=28, color=MUTED, spacing=1.3)
footer(s, "18 · спасибо! вопросы · [контакты докладчика]")

order = json.load(open(os.path.join(HERE, "deck.json")))["order"]
assert len(prs.slides) == len(order), (len(prs.slides), len(order))
out = os.path.join(HERE, "talk.pptx")
prs.save(out)
print(out, len(prs.slides), "slides")
